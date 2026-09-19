"""
Controlador del Instructor — Gestión de sesiones QR, historial y alertas.
"""
from flask import Blueprint, request, render_template, redirect, url_for, session, jsonify, send_file, flash
from datetime import datetime, date
from controllers.auth_controller import requiere_rol
from models import ficha, sesion as modelo_sesion, asistencia, alerta, auditoria
from models.dashboard import obtener_metricas_instructor
from controllers.seguridad import generar_codigo_qr, obtener_ip_real
from config import Config
import qrcode
import io
import base64
import socket
import os

instructor_bp = Blueprint('instructor', __name__)


def obtener_host_publico(req):
    """
    Retorna la URL base adecuada para el QR y para compartir con los aprendices:
    1. Si existe SERVER_PUBLIC_URL en variables de entorno, la usa.
    2. Si no es localhost (ej: dominio en Render o túnel Cloudflare), usa el host actual.
    3. Si es localhost, detecta la IP de red local (ej: 192.168.40.13:5000) para que
       cualquier celular conectado al mismo Wi-Fi pueda abrirlo al escanear.
    """
    env_url = os.environ.get('SERVER_PUBLIC_URL')
    if env_url:
        return env_url.rstrip('/')

    host = req.host.lower()
    if 'localhost' in host or '127.0.0.1' in host:
        try:
            s = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
            s.connect(("8.8.8.8", 80))
            ip = s.getsockname()[0]
            s.close()
        except Exception:
            ip = '127.0.0.1'
        port = f":{req.host.split(':')[1]}" if ':' in req.host else ':5000'
        return f"http://{ip}{port}"

    scheme = req.headers.get('X-Forwarded-Proto', req.scheme)
    return f"{scheme}://{req.host}"


def generar_imagen_qr(contenido):
    """Genera una imagen QR en formato base64 para mostrar en el navegador."""
    qr = qrcode.QRCode(version=1, box_size=10, border=2)
    qr.add_data(contenido)
    qr.make(fit=True)
    img = qr.make_image(fill_color='#12140F', back_color='white')
    buffer = io.BytesIO()
    img.save(buffer, format='PNG')
    buffer.seek(0)
    return base64.b64encode(buffer.getvalue()).decode('utf-8')


# ============================================================
#  SESIÓN DE ASISTENCIA (Pantalla 1)
# ============================================================

@instructor_bp.route('/instructor/sesion')
@requiere_rol('instructor')
def sesion():
    """Muestra la pantalla de sesión con QR activo o selector para iniciar."""
    # 1. Obtener fichas activas asignadas a este instructor
    fichas_asignadas = ficha.obtener_por_instructor(session['user_id'], estado='activa') or []
    
    # 2. Obtener todas las fichas activas del centro de formación
    todas_activas = ficha.obtener_todas(estado='activa') or []
    
    # Priorizar fichas asignadas al instructor, luego las demás activas sin duplicar
    ids_vistos = set()
    fichas = []
    for f in fichas_asignadas + todas_activas:
        if f['id'] not in ids_vistos:
            ids_vistos.add(f['id'])
            fichas.append(f)
            
    # Fallback si ninguna está activa
    if not fichas:
        fichas = ficha.obtener_todas() or []

    activa = modelo_sesion.obtener_activa_por_instructor(session['user_id'])
    metricas = obtener_metricas_instructor(session['user_id'])
    url_compartir = obtener_host_publico(request)
    return render_template('instructor/sesion.html', fichas=fichas, sesion_activa=activa, metricas=metricas, url_compartir=url_compartir)


@instructor_bp.route('/instructor/sesion/generar', methods=['POST'])
@requiere_rol('instructor')
def generar_sesion():
    """Crea una nueva sesión de asistencia y genera el primer QR."""
    ficha_id = request.form.get('ficha_id')
    if not ficha_id:
        flash('Debes seleccionar una ficha válida para iniciar la sesión', 'warning')
        return redirect(url_for('instructor.sesion'))
        
    # Validar si ya tiene una sesión activa
    activa = modelo_sesion.obtener_activa_por_instructor(session['user_id'])
    if activa:
        flash('Ya tienes una sesión activa en curso para esta u otra ficha', 'warning')
        return redirect(url_for('instructor.sesion'))
        
    ip_instructor = obtener_ip_real(request)
    
    # Crear la sesión (guarda la IP del instructor para validación de red)
    sesion_id = modelo_sesion.crear(ficha_id, session['user_id'], ip_instructor)
    
    # Generar el primer código QR firmado con HMAC
    codigo = generar_codigo_qr(sesion_id, Config.QR_HMAC_SECRET)
    modelo_sesion.actualizar_codigo(sesion_id, codigo)
    
    # Registrar en auditoría
    auditoria.registrar(session['user_id'], ip_instructor, 'SESION_CREADA', 'exitoso', f'Sesión {sesion_id}, Ficha {ficha_id}')
    
    return redirect(url_for('instructor.sesion'))


@instructor_bp.route('/instructor/sesion/cerrar', methods=['POST'])
@requiere_rol('instructor')
def cerrar_sesion():
    """Cierra la ventana de registro de una sesión y marca inasistencias de forma atómica."""
    sesion_id_form = request.form.get('sesion_id')
    
    if sesion_id_form:
        sesion_target = modelo_sesion.obtener_por_id(sesion_id_form)
        if sesion_target and (sesion_target['instructor_id'] == session['user_id'] or session.get('rol') == 'admin'):
            modelo_sesion.cerrar(sesion_id_form)
            ip = obtener_ip_real(request)
            auditoria.registrar(session['user_id'], ip, 'SESION_CERRADA', 'exitoso', f'Sesión {sesion_id_form}')
            flash('Sesión de asistencia cerrada correctamente. Se consolidaron las ausencias.', 'success')
    else:
        activa = modelo_sesion.obtener_activa_por_instructor(session['user_id'])
        if activa:
            modelo_sesion.cerrar(activa['id'])
            ip = obtener_ip_real(request)
            auditoria.registrar(session['user_id'], ip, 'SESION_CERRADA', 'exitoso', f'Sesión {activa["id"]}')
            flash('Ventana de asistencia cerrada. Se marcaron ausencias automáticamente.', 'success')
            
    redirect_url = request.referrer or url_for('instructor.sesion')
    return redirect(redirect_url)


@instructor_bp.route('/instructor/sesion/qr')
@requiere_rol('instructor')
def api_qr():
    """API: Genera un nuevo código QR (se llama cada 30 segundos via AJAX)."""
    activa = modelo_sesion.obtener_activa_por_instructor(session['user_id'])
    if not activa:
        return jsonify({'error': 'No hay sesión activa'}), 400
    
    # Generar nuevo código HMAC
    codigo = generar_codigo_qr(activa['id'], Config.QR_HMAC_SECRET)
    modelo_sesion.actualizar_codigo(activa['id'], codigo)
    
    # Generar imagen QR como base64 con la URL accesible para cualquier teléfono
    base_url = obtener_host_publico(request)
    ruta_aprendiz = url_for('aprendiz.escanear', codigo=codigo)
    url_qr = f"{base_url}{ruta_aprendiz}"
    imagen_base64 = generar_imagen_qr(url_qr)
    
    return jsonify({
        'codigo': codigo,
        'imagen_base64': imagen_base64,
        'url_qr': url_qr,
        'base_url': base_url
    })


@instructor_bp.route('/instructor/sesion/registros')
@requiere_rol('instructor')
def registros_sesion():
    """API: Retorna la lista de asistentes registrados en la sesión activa (polling cada 5s)."""
    activa = modelo_sesion.obtener_activa_por_instructor(session['user_id'])
    if not activa:
        return jsonify({'registros': [], 'total_aprendices': 0})
    
    # obtener_por_sesion() ya retorna el dict formateado para JSON
    datos = asistencia.obtener_por_sesion(activa['id'])
    
    # Agregar el total de aprendices de la ficha
    aprendices = ficha.obtener_aprendices(activa['ficha_id'])
    datos['total_aprendices'] = len(aprendices) if aprendices else 0
    
    return jsonify(datos)


# ============================================================
#  HISTORIAL DE ASISTENCIA (Pantalla 3)
# ============================================================

@instructor_bp.route('/instructor/historial')
@requiere_rol('instructor')
def historial():
    """Muestra el historial de asistencia con filtros, métricas y paginación real."""
    ficha_id = request.args.get('ficha_id') or request.args.get('ficha')
    fecha_inicio = request.args.get('fecha_inicio')
    fecha_fin = request.args.get('fecha_fin')
    estado = request.args.get('estado')
    busqueda = request.args.get('q') or request.args.get('busqueda')
    
    try:
        pagina = int(request.args.get('pagina', 1))
        if pagina < 1:
            pagina = 1
    except (ValueError, TypeError):
        pagina = 1
        
    por_pagina = 15
    
    # Obtener historial filtrado
    hist = asistencia.obtener_historial(
        session['user_id'], ficha_id, fecha_inicio, fecha_fin, estado, busqueda, pagina, por_pagina=por_pagina
    )
    total = hist.get('total', 0)
    total_paginas = max(1, (total + por_pagina - 1) // por_pagina)
    
    # Calcular métricas sincronizadas con los filtros activos
    metricas = asistencia.calcular_metricas(session['user_id'], ficha_id, fecha_inicio, fecha_fin)
    
    # Obtener fichas disponibles para el selector de filtro
    fichas_asignadas = ficha.obtener_por_instructor(session['user_id']) or []
    todas_fichas = ficha.obtener_todas() or []
    ids_vistos = set()
    fichas = []
    for f in fichas_asignadas + todas_fichas:
        if f['id'] not in ids_vistos:
            ids_vistos.add(f['id'])
            fichas.append(f)
    
    # Obtener sesiones anteriores del instructor
    sesiones = modelo_sesion.obtener_todas_por_instructor(session['user_id'])
    
    return render_template('instructor/historial.html',
        registros=hist['registros'],
        total=total,
        total_paginas=total_paginas,
        pagina=pagina,
        por_pagina=por_pagina,
        metricas=metricas,
        fichas=fichas,
        sesiones=sesiones,
        filtros={
            'ficha': ficha_id or '',
            'ficha_id': ficha_id or '',
            'fecha_inicio': fecha_inicio or '',
            'fecha_fin': fecha_fin or '',
            'estado': estado or '',
            'busqueda': busqueda or '',
            'q': busqueda or ''
        }
    )


@instructor_bp.route('/instructor/historial/corregir', methods=['POST'])
@requiere_rol('instructor')
def corregir_asistencia():
    """Corrige manualmente un registro de asistencia con auditoría completa."""
    asistencia_id = request.form.get('asistencia_id')
    nuevo_estado = request.form.get('estado', 'corregido')
    motivo = (request.form.get('motivo') or '').strip()
    
    if not asistencia_id:
        flash('Registro de asistencia no especificado.', 'error')
        return redirect(url_for('instructor.historial'))
        
    if not motivo:
        motivo = 'Ajuste manual de asistencia realizado por el instructor.'
    
    asistencia.corregir(asistencia_id, nuevo_estado, motivo)
    
    ip = obtener_ip_real(request)
    auditoria.registrar(
        session['user_id'],
        ip,
        'CORRECCION_ASISTENCIA',
        'exitoso',
        f'Registro ID {asistencia_id} modificado a estado "{nuevo_estado}". Justificación: {motivo}'
    )
    flash(f'Registro actualizado correctamente a "{nuevo_estado}".', 'success')
    
    # Retornar a la misma página y filtros
    redirect_url = request.referrer or url_for('instructor.historial')
    return redirect(redirect_url)


@instructor_bp.route('/instructor/historial/exportar')
@requiere_rol('instructor')
def exportar_excel():
    """Genera un archivo Excel ejecutivo con el historial filtrado y diseño institucional SENA."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    
    ficha_id = request.args.get('ficha_id') or request.args.get('ficha')
    fecha_inicio = request.args.get('fecha_inicio')
    fecha_fin = request.args.get('fecha_fin')
    estado = request.args.get('estado')
    busqueda = request.args.get('q') or request.args.get('busqueda')
    
    # Obtener todo el conjunto de datos que coincide con el filtro
    hist = asistencia.obtener_historial(
        session['user_id'], ficha_id, fecha_inicio, fecha_fin, estado, busqueda, pagina=1, por_pagina=100000
    )
    registros = hist.get('registros', [])
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Historial Asistencia"
    ws.views.sheetView[0].showGridLines = True
    
    # Colores corporativos SENA
    COLOR_SENA_VERDE = "39A900"
    COLOR_SENA_VERDE_OSCURO = "287800"
    COLOR_CABECERA_TABLA = "1E293B"
    COLOR_PRESENTE = "D1E7DD"
    COLOR_AUSENTE = "F8D7DA"
    COLOR_CORREGIDO = "FFF3CD"
    
    # Estilos de bordes
    borde_fino = Border(
        left=Side(style='thin', color='E2E8F0'),
        right=Side(style='thin', color='E2E8F0'),
        top=Side(style='thin', color='E2E8F0'),
        bottom=Side(style='thin', color='E2E8F0')
    )
    borde_cabecera = Border(
        left=Side(style='thin', color='334155'),
        right=Side(style='thin', color='334155'),
        top=Side(style='medium', color='1E293B'),
        bottom=Side(style='medium', color='1E293B')
    )
    
    # Encabezado Institucional
    ws.merge_cells('A1:I1')
    ws['A1'] = "SERVICIO NACIONAL DE APRENDIZAJE - SENA"
    ws['A1'].font = Font(name='Calibri', size=15, bold=True, color='FFFFFF')
    ws['A1'].fill = PatternFill(start_color=COLOR_SENA_VERDE, end_color=COLOR_SENA_VERDE, fill_type='solid')
    ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[1].height = 32
    
    ws.merge_cells('A2:I2')
    ws['A2'] = "CONTROL Y REGISTRO BIOMÉTRICO/QR DE ASISTENCIA — REPORTE DETALLADO"
    ws['A2'].font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
    ws['A2'].fill = PatternFill(start_color=COLOR_SENA_VERDE_OSCURO, end_color=COLOR_SENA_VERDE_OSCURO, fill_type='solid')
    ws['A2'].alignment = Alignment(horizontal='center', vertical='center')
    ws.row_dimensions[2].height = 22
    
    # Metadatos del Reporte
    ws['A4'] = "Generado por:"
    ws['A4'].font = Font(bold=True, color='475569')
    ws['B4'] = session.get('nombre', 'Instructor SENA')
    ws['B4'].font = Font(bold=True)
    
    ws['E4'] = "Fecha y hora:"
    ws['E4'].font = Font(bold=True, color='475569')
    ws['F4'] = datetime.now().strftime('%d/%m/%Y %H:%M:%S')
    
    ws['A5'] = "Total registros:"
    ws['A5'].font = Font(bold=True, color='475569')
    ws['B5'] = len(registros)
    ws['B5'].font = Font(bold=True)
    
    filtro_desc = []
    if ficha_id:
        filtro_desc.append(f"Ficha ID: {ficha_id}")
    if estado:
        filtro_desc.append(f"Estado: {estado.capitalize()}")
    if fecha_inicio or fecha_fin:
        filtro_desc.append(f"Período: {fecha_inicio or 'Inicio'} al {fecha_fin or 'Hoy'}")
    if busqueda:
        filtro_desc.append(f'Búsqueda: "{busqueda}"')
        
    ws['E5'] = "Filtros aplicados:"
    ws['E5'].font = Font(bold=True, color='475569')
    ws['F5'] = " | ".join(filtro_desc) if filtro_desc else "Todos los registros (Sin filtros)"
    
    # Encabezados de la Tabla (Fila 7)
    headers = [
        "Ficha", "Programa de Formación", "Documento", "Nombre del Aprendiz",
        "Fecha", "Hora", "Estado", "Método", "Motivo de Corrección"
    ]
    ws.row_dimensions[7].height = 24
    
    for col_idx, header in enumerate(headers, start=1):
        cell = ws.cell(row=7, column=col_idx, value=header)
        cell.font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
        cell.fill = PatternFill(start_color=COLOR_CABECERA_TABLA, end_color=COLOR_CABECERA_TABLA, fill_type='solid')
        cell.alignment = Alignment(horizontal='center', vertical='center')
        cell.border = borde_cabecera
        
    # Datos de filas (Fila 8+)
    row_idx = 8
    for r in registros:
        ws.row_dimensions[row_idx].height = 20
        
        fecha_val = r.get('fecha_formateada') or str(r.get('fecha') or '')
        hora_val = r.get('hora_formateada') or str(r.get('hora_registro') or '')
        estado_raw = (r.get('estado') or '').lower()
        estado_fmt = estado_raw.capitalize() if estado_raw else 'Presente'
        
        fila_datos = [
            r.get('ficha_numero', ''),
            r.get('ficha_programa') or r.get('programa', ''),
            r.get('documento', ''),
            r.get('nombre', ''),
            fecha_val,
            hora_val,
            estado_fmt,
            (r.get('metodo') or 'qr').upper(),
            r.get('motivo_correccion', '') or ''
        ]
        
        for col_idx, val in enumerate(fila_datos, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = Font(name='Calibri', size=10)
            cell.border = borde_fino
            
            # Alineaciones específicas
            if col_idx in (1, 3, 5, 6, 7, 8):
                cell.alignment = Alignment(horizontal='center', vertical='center')
            else:
                cell.alignment = Alignment(horizontal='left', vertical='center')
                
            # Formato condicional de Estado
            if col_idx == 7:
                if estado_raw == 'presente':
                    cell.fill = PatternFill(start_color=COLOR_PRESENTE, end_color=COLOR_PRESENTE, fill_type='solid')
                    cell.font = Font(name='Calibri', size=10, bold=True, color='0F5132')
                elif estado_raw == 'ausente':
                    cell.fill = PatternFill(start_color=COLOR_AUSENTE, end_color=COLOR_AUSENTE, fill_type='solid')
                    cell.font = Font(name='Calibri', size=10, bold=True, color='842029')
                elif estado_raw == 'corregido':
                    cell.fill = PatternFill(start_color=COLOR_CORREGIDO, end_color=COLOR_CORREGIDO, fill_type='solid')
                    cell.font = Font(name='Calibri', size=10, bold=True, color='664D03')
            elif row_idx % 2 == 1:
                # Zebra striping suave
                cell.fill = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
                
        row_idx += 1
        
    # Auto-ajuste de ancho de columnas
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            # Omitir filas de título fusionadas para el cálculo de ancho
            if cell.row in (1, 2, 4, 5):
                continue
            if cell.value is not None:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)
        
    # Ancho específico para nombres y programas
    ws.column_dimensions['B'].width = max(ws.column_dimensions['B'].width, 26)
    ws.column_dimensions['D'].width = max(ws.column_dimensions['D'].width, 30)
    ws.column_dimensions['I'].width = max(ws.column_dimensions['I'].width, 28)
    
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    
    nombre_archivo = f"Asistencia_SENA_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return send_file(
        out,
        mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        as_attachment=True,
        download_name=nombre_archivo
    )

# ============================================================
#  ALERTAS DE DESERCIÓN (Pantalla 4)
# ============================================================

@instructor_bp.route('/instructor/alertas')
@requiere_rol('instructor')
def alertas():
    """Muestra las alertas tempranas de deserción."""
    alertas_lista = alerta.obtener_por_instructor(session['user_id'])
    metricas_alertas = alerta.contar_activas(session['user_id'])
    return render_template('instructor/alertas.html', alertas=alertas_lista, metricas=metricas_alertas)


@instructor_bp.route('/instructor/alertas/atender', methods=['POST'])
@requiere_rol('instructor')
def atender_alerta():
    """Marca una alerta como atendida con nota de seguimiento."""
    alerta_id = request.form.get('alerta_id')
    nota = request.form.get('nota', '')
    
    alerta.atender(alerta_id, nota)
    
    ip = obtener_ip_real(request)
    auditoria.registrar(session['user_id'], ip, 'ALERTA_ATENDIDA', 'exitoso', f'Alerta {alerta_id}')
    
    return redirect(url_for('instructor.alertas'))
